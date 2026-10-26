"""Geração de documentos .docx (modelos) com mesclagem de dados do cliente/processo.

Modelos disponíveis: procuração ad judicia e contrato de honorários.
São modelos genéricos — o advogado pode editar o .docx gerado livremente.
"""
import io
from datetime import datetime

MESES = ["", "janeiro", "fevereiro", "março", "abril", "maio", "junho",
         "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]

MODELOS = {
    "procuracao": "Procuração",
    "contrato": "Contrato de Honorários",
}


def _hoje_extenso() -> str:
    d = datetime.now()
    return f"Maceió, {d.day} de {MESES[d.month]} de {d.year}."


def _save(doc) -> io.BytesIO:
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def gerar_modelo(modelo: str, cliente, processo=None) -> io.BytesIO:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(12)

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("LEOJURIS")
    run.bold = True
    run.font.size = Pt(14)

    nome = cliente.name or "—"
    doc_cli = cliente.cpf_cnpj or "—"
    proc_num = (processo.case_number if processo else None) or "—"

    if modelo == "contrato":
        _contrato(doc, nome, doc_cli, cliente, proc_num)
    else:
        _procuracao(doc, nome, doc_cli, cliente, proc_num)

    return _save(doc)


def _procuracao(doc, nome, doc_cli, cliente, proc_num):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("\nPROCURAÇÃO AD JUDICIA ET EXTRA")
    r.bold = True
    r.font.size = Pt(13)

    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("OUTORGANTE: ").bold = True
    contato = []
    if cliente.email:
        contato.append(f"e-mail {cliente.email}")
    if cliente.phone:
        contato.append(f"telefone {cliente.phone}")
    extra = (", " + ", ".join(contato)) if contato else ""
    p.add_run(
        f"{nome}, inscrito(a) no CPF/CNPJ sob o nº {doc_cli}{extra}, "
        f"doravante denominado(a) OUTORGANTE."
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("OUTORGADO: ").bold = True
    p.add_run(
        "os advogados do escritório LeoJuris, regularmente inscritos "
        "na OAB, com escritório profissional onde recebem intimações."
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "PODERES: pelo presente instrumento, o(a) outorgante nomeia e constitui seus "
        "bastantes procuradores os outorgados acima, a quem confere os poderes da "
        "cláusula ad judicia et extra, para o foro em geral, em qualquer Juízo, "
        "Instância ou Tribunal, podendo propor contra quem de direito as ações "
        "competentes e defendê-lo(a) nas contrárias, seguindo umas e outras até final "
        "decisão, usando os recursos legais e acompanhando-os, conferindo ainda poderes "
        "especiais para confessar, desistir, transigir, firmar compromissos ou acordos, "
        "receber e dar quitação, agindo em conjunto ou separadamente."
    )

    if proc_num and proc_num != "—":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(f"Processo de referência: {proc_num}.")

    doc.add_paragraph("")
    data = doc.add_paragraph(_hoje_extenso())
    data.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph("\n\n")
    ass = doc.add_paragraph("_______________________________________")
    ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nm = doc.add_paragraph(nome)
    nm.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _contrato(doc, nome, doc_cli, cliente, proc_num):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("\nCONTRATO DE PRESTAÇÃO DE SERVIÇOS ADVOCATÍCIOS")
    r.bold = True
    r.font.size = Pt(13)

    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("CONTRATANTE: ").bold = True
    p.add_run(f"{nome}, inscrito(a) no CPF/CNPJ sob o nº {doc_cli}.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("CONTRATADO: ").bold = True
    p.add_run("LeoJuris, por seus advogados regularmente inscritos na OAB.")

    for titulo, texto in [
        ("CLÁUSULA 1ª - DO OBJETO",
         "O presente contrato tem por objeto a prestação de serviços advocatícios pelo "
         "CONTRATADO ao CONTRATANTE" + (f", referente ao processo {proc_num}" if proc_num and proc_num != "—" else "") + "."),
        ("CLÁUSULA 2ª - DOS HONORÁRIOS",
         "Pelos serviços prestados, o CONTRATANTE pagará ao CONTRATADO honorários no valor "
         "de R$ __________ (______________________), na forma e nas datas ajustadas entre as partes."),
        ("CLÁUSULA 3ª - DAS OBRIGAÇÕES",
         "O CONTRATADO obriga-se a empregar zelo e diligência no patrocínio da causa, e o "
         "CONTRATANTE a fornecer os documentos e informações necessários ao bom andamento do trabalho."),
        ("CLÁUSULA 4ª - DO FORO",
         "Fica eleito o foro da comarca de Maceió/AL para dirimir quaisquer questões oriundas do presente contrato."),
    ]:
        ct = doc.add_paragraph()
        ct.add_run(titulo).bold = True
        cp = doc.add_paragraph(texto)
        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph("")
    data = doc.add_paragraph(_hoje_extenso())
    data.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph("\n")
    for label in ("CONTRATANTE", "CONTRATADO"):
        a = doc.add_paragraph("_______________________________________")
        a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        l = doc.add_paragraph(label)
        l.alignment = WD_ALIGN_PARAGRAPH.CENTER
