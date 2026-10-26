"""Geração de planilhas (.xlsx) e tabelas em PDF para exportação de listas."""
import io

# 🚨 CSV/Formula injection: uma célula que COMEÇA com = + - @ (ou tab/quebra) é
# executada pelo Excel/LibreOffice ao abrir — o vetor clássico de =WEBSERVICE()
# e execução via DDE. Pior no openpyxl: um texto começando com "=" é gravado
# como FÓRMULA VIVA (<f>), não como texto (confirmado em pentest, ago/2026).
# Como QUALQUER usuário digita o nome/observação que vai para a planilha, e quem
# abre o export é o admin, o dado do usuário não pode virar fórmula na mão dele.
_GATILHOS_FORMULA = ("=", "+", "-", "@", "\t", "\r", "\n")


def _escrever_linha(ws, valores):
    """Escreve uma linha neutralizando célula que pareça fórmula.

    A neutralização preserva o VALOR (telefone "+55 81..." continua "+55 81...",
    sem apóstrofo no dado): força o tipo texto e marca `quotePrefix`, que faz o
    Excel exibir como texto puro, sem marcador visível. Valor normal passa
    intacto pelo caminho de sempre.
    """
    linha = ws.max_row + 1
    for col, v in enumerate(valores, 1):
        cell = ws.cell(row=linha, column=col)
        if isinstance(v, str) and v[:1] in _GATILHOS_FORMULA:
            cell.value = v
            cell.data_type = "s"      # nunca fórmula: guarda como string
            cell.quotePrefix = True   # Excel mostra como texto, sem o apóstrofo
        else:
            cell.value = "" if v is None else v


def to_xlsx(headers, rows, title="Dados") -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Dados")[:31]

    head_fill = PatternFill("solid", fgColor="1A3A5C")
    head_font = Font(color="FFFFFF", bold=True)
    ws.append(list(headers))
    for i in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=i)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        _escrever_linha(ws, r)

    # Largura aproximada das colunas
    for i, h in enumerate(headers, 1):
        col_vals = [len(str(h))] + [len(str(r[i - 1])) for r in rows] if rows else [len(str(h))]
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(col_vals) + 2, 48)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def to_pdf_table(headers, rows, title="Relatório") -> io.BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=1 * cm, rightMargin=1 * cm, topMargin=1 * cm, bottomMargin=1 * cm,
    )
    styles = getSampleStyleSheet()
    elems = [
        Paragraph(f"<b>LeoJuris</b> — {title}", styles["Title"]),
        Spacer(1, 10),
    ]
    data = [list(headers)] + [["" if v is None else str(v) for v in r] for r in rows]
    if not rows:
        data.append(["(sem registros)"] + [""] * (len(headers) - 1))
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A3A5C")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elems.append(t)
    doc.build(elems)
    buf.seek(0)
    return buf
